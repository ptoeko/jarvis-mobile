"""Сквозные тесты версии 0.4.1: экспорт реестра в Excel (.xlsx).

Запуск: python3 tests/e2e/test_v041.py [папка]

Файл скачивается из приложения (Chromium, экран Pixel 7), открывается
openpyxl и LibreOffice; проверяется каждая ячейка: значения как в реестре
(отметки, оси, даты не превращаются в числа), «Не указано», переносы строк,
«формула» остаётся текстом, имена файлов в последнем столбце; закреплённая
шапка и фильтр. В варианте APK — сохранение через окно Android (имитация моста).
Нужны: playwright с Chromium, openpyxl, LibreOffice (soffice).
"""
import csv
import io
import subprocess
import sys
import traceback
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_stage2 as base  # noqa: E402
import test_v040 as v040  # noqa: E402
import openpyxl  # noqa: E402
from playwright.sync_api import sync_playwright, expect  # noqa: E402

OUT = base.OUT
NS = base.NS
TITLES = base.TITLES + ['Файлы (вложения)']

NOTE1 = ('23.10.2026. Объект: корпус 7. Ось 3/Б, на отметке -0,300 — раковины "до 5 мм"\n'
         'вторая строка. Подрядчик: ООО «Альфа, Бета и К», срок до 30.10.2026')
NOTE2 = '=СУММ(A1:A9) отм. +12,450 — трещина в стяжке'


def make_records(page):
    v040.to_dialog(page)
    v040.attach(page, 'фото трещины.png', 'акт освидетельствования.pdf')
    page.fill('#input', f'{NOTE1}\n\n{NOTE2}')
    page.click('#parse')
    expect(page.locator('.draft')).to_have_count(2)
    base.draft(page, 0).get_by_role('button', name='Добавить в реестр').click()
    expect(page.locator('.draft')).to_have_count(1)
    base.draft(page, 0).get_by_role('button', name='Добавить в реестр').click()
    expect(page.locator('.record-row')).to_have_count(2)


def expected_rows(page):
    """Значения из реестра приложения — эталон для проверки файла."""
    return page.evaluate('''async () => {
        const db = await new Promise(r => { const q = indexedDB.open('jarvis-mobile'); q.onsuccess = () => r(q.result); });
        const recs = await new Promise(r => { const q = db.transaction('records').objectStore('records').getAll(); q.onsuccess = () => r(q.result); });
        db.close();
        const keys = ['number','date','object','location','workType','description','requirement','responsible','deadline','status','sourceText'];
        return recs.sort((a, b) => a.number - b.number).map(r => [...keys.map(k => r[k]), (r.attachments || []).map(a => a.name).join('\\n')]);
    }''')


def check_workbook(path, expected):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.title == 'Реестр', ws.title
    assert ws.max_column == 12 and ws.max_row == len(expected) + 1, (ws.max_column, ws.max_row)
    assert [c.value for c in ws[1]] == TITLES, [c.value for c in ws[1]]
    assert ws.freeze_panes == 'A2', ws.freeze_panes
    assert ws.auto_filter.ref == f'A1:L{len(expected) + 1}', ws.auto_filter.ref
    assert all(c.font.b for c in ws[1]), 'шапка жирная'
    for r, row in enumerate(expected, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c)
            if c == 1:
                assert cell.value == value and isinstance(cell.value, int), (r, c, cell.value)
            else:
                assert cell.data_type == 's', (r, c, cell.data_type, cell.value)
                assert (cell.value or '') == value, (r, c, cell.value, value)
    return ws


def libreoffice_values(path):
    out = path.parent / 'lo'
    out.mkdir(exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', 'csv:Text - txt - csv (StarCalc):44,34,76',
                    '--outdir', str(out), str(path)], check=True, capture_output=True, timeout=180)
    text = (out / (path.stem + '.csv')).read_text(encoding='utf-8')
    return list(csv.reader(io.StringIO(text)))


TESTS = []


def case(fn):
    TESTS.append(fn)
    return fn


@case
def export_xlsx_every_cell(page, ctx):
    make_records(page)
    expect(page.locator('#export-xlsx')).to_be_enabled()
    page.click('#export-xlsx')
    expect(page.locator('#export-xlsx-summary')).to_contain_text('записей: 2, столбцов: 12')
    base.shot(page, 'v041-01-export-confirm')
    with page.expect_download() as dl:
        page.click('#export-xlsx-confirm')
    name = dl.value.suggested_filename
    assert name.startswith('jarvis-reestr-') and name.endswith('.xlsx'), name
    path = OUT / name
    dl.value.save_as(path)
    expect(page.locator('#toast')).to_contain_text('.xlsx')

    expected = expected_rows(page)
    assert expected[0][3] == 'Ось 3/Б, на отметке -0,300', expected[0][3]
    assert expected[0][11] == 'фото трещины.png\nакт освидетельствования.pdf'
    assert expected[1][9] == NS
    check_workbook(path, expected)

    # LibreOffice открывает файл и видит те же значения (без превращения в числа).
    rows = libreoffice_values(path)
    assert rows[0] == TITLES, rows[0]
    for got, want in zip(rows[1:], expected):
        assert got == [str(want[0])] + want[1:], (got, want)
    assert any(v.startswith('=СУММ') for v in rows[2]), 'формула осталась текстом'


@case
def registry_buttons_fit_phone(page, ctx):
    page.set_viewport_size({'width': 360, 'height': 740})
    make_records(page)
    overflow = page.evaluate('document.documentElement.scrollWidth - document.documentElement.clientWidth')
    assert overflow <= 0, f'горизонтальная прокрутка {overflow}px'
    for sel in ['#export-xlsx', '#export']:
        box = page.locator(sel).bounding_box()
        assert box['height'] >= 44 and box['x'] >= 0 and box['x'] + box['width'] <= 360, (sel, box)
    page.evaluate("document.getElementById('main').scrollTop = 0")
    base.shot(page, 'v041-02-registry-360px')


@case
def export_disabled_when_empty(page, ctx):
    page.click('.tab[data-target="registry"]')
    expect(page.locator('#export-xlsx')).to_be_disabled()
    expect(page.locator('#export')).to_be_disabled()


def apk_export(page, ctx):
    make_records(page)
    page.click('#export-xlsx')
    expect(page.locator('#export-xlsx-confirm')).to_have_text('Сохранить Excel')
    page.click('#export-xlsx-confirm')
    expect(page.locator('#toast')).to_contain_text('сохранён')
    saved = page.evaluate('window.__saved')
    assert len(saved) == 1 and saved[0]['mime'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', saved
    data = page.evaluate('(t) => window.__chunks[t]', saved[0]['token']).encode('latin-1')
    path = OUT / 'apk-export.xlsx'
    path.write_bytes(data)
    assert zipfile.is_zipfile(path)
    check_workbook(path, expected_rows(page))


def main():
    srv = base.Server(base.ROOT)
    v040.srv_url = srv.url
    failures = 0
    with sync_playwright() as p:
        browser = p.chromium.launch()
        device = p.devices['Pixel 7']
        failures += v040.run(TESTS, browser, device, 'web')
        failures += v040.run([apk_export], browser, device, 'apk')
        browser.close()
    srv.httpd.shutdown()
    total = len(TESTS) + 1
    print(f'\n{total - failures} из {total} тестов пройдено. Файлы: {OUT}')
    sys.exit(1 if failures else 0)


if __name__ == '__main__':
    main()
